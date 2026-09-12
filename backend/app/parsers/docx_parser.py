from __future__ import annotations

import re
from collections import defaultdict
from datetime import time
from pathlib import Path
from urllib.parse import urlparse

from docx import Document
from docx.table import _Cell, _Row, Table

from app.parsers.base import ParseResult, ParsedLesson

DAY_NAMES = {
    "ПОНЕДІЛОК": 0,
    "ВІВТОРОК": 1,
    "СЕРЕДА": 2,
    "ЧЕТВЕР": 3,
    "П’ЯТНИЦЯ": 4,
    "П'ЯТНИЦЯ": 4,
}
LESSON_TIMES = {
    1: (time(8, 30), time(10, 0)),
    2: (time(10, 15), time(11, 45)),
    3: (time(12, 15), time(13, 45)),
    4: (time(14, 0), time(15, 30)),
}
URL_RE = re.compile(r"https?://\S+", re.IGNORECASE)
WEEK_PREFIX_RE = re.compile(r"^\s*(?P<start>\d{1,2})\s*-\s*(?P<end>\d{1,2})\s*(?P<rest>.*)$", re.IGNORECASE)
TYPE_WORDS = {"лекція", "лекция", "лаб.", "лаб", "практика", "практ."}


class DocxScheduleParser:
    """Parser for the actual provided Word schedule: weekly grid, merged day/group cells."""

    def parse(self, path: Path) -> ParseResult:
        result = ParseResult()
        doc = Document(path)
        if not doc.tables:
            result.errors.append("DOCX does not contain schedule tables.")
            return result
        seen: set[tuple] = set()
        known_group_map: dict[int, str] = {}
        for table_index, table in enumerate(doc.tables):
            group_map = self._group_map(table, known_group_map)
            if group_map:
                known_group_map = group_map
            current_day: int | None = None
            current_lesson_number: int | None = None
            for row_index, row in enumerate(table.rows):
                expanded = self._expanded_row(row)
                if self._is_group_header(expanded):
                    continue
                if len(expanded) < 3:
                    continue
                day_text = self._compact(expanded[0][0])
                if day_text in DAY_NAMES:
                    current_day = DAY_NAMES[day_text]
                lesson_text = self._compact(expanded[1][0])
                if lesson_text.isdigit():
                    current_lesson_number = int(lesson_text)
                if current_lesson_number is None or current_day is None:
                    continue
                lesson_number = current_lesson_number
                result.total_rows += 1
                for col_index, (raw_text, _cell_id) in enumerate(expanded[2:], start=2):
                    text = self._normalize_text(raw_text)
                    if not text or self._is_elective(text):
                        continue
                    groups = self._groups_for_span(expanded, col_index, group_map)
                    if not groups:
                        result.warnings.append(f"Table {table_index + 1}, row {row_index + 1}: cannot infer group for column {col_index + 1}.")
                        continue
                    parsed = self._parse_cell(text, lesson_number)
                    for group_name in groups:
                        lesson = ParsedLesson(
                            group_name=group_name,
                            day_of_week=current_day,
                            lesson_number=lesson_number,
                            starts_at=parsed["starts_at"],
                            ends_at=parsed["ends_at"],
                            subject=parsed["subject"],
                            teacher=parsed["teacher"],
                            meeting_url=parsed["meeting_url"],
                            room=None,
                            description=parsed["description"],
                            lesson_type=parsed["lesson_type"],
                            week_start=parsed["week_start"],
                            week_end=parsed["week_end"],
                            week_parity=parsed["week_parity"],
                            source_table=table_index + 1,
                            source_row=row_index + 1,
                            source_col=col_index + 1,
                        )
                        key = (
                            lesson.group_name,
                            lesson.day_of_week,
                            lesson.lesson_number,
                            lesson.subject,
                            lesson.teacher,
                            lesson.meeting_url,
                            lesson.week_start,
                            lesson.week_end,
                            lesson.week_parity,
                        )
                        if key not in seen:
                            seen.add(key)
                            result.lessons.append(lesson)
        result.warnings.append(
            "DOCX contains a weekly recurring grid without calendar dates; dates are derived from SEMESTER_START_DATE."
        )
        result.warnings.append("DOCX does not contain explicit lesson times; configured default lesson times are used.")
        return self._validate(result)

    def _expanded_row(self, row: _Row) -> list[tuple[str, int]]:
        expanded: list[tuple[str, int]] = []
        for raw_index, tc in enumerate(row._tr.tc_lst):
            cell = _Cell(tc, row.table)
            span = int(cell._tc.tcPr.gridSpan.val) if cell._tc.tcPr is not None and cell._tc.tcPr.gridSpan is not None else 1
            text = cell.text
            cell_id = id(tc)
            expanded.extend((text, cell_id) for _ in range(span))
        return expanded

    def _group_map(self, table: Table, previous: dict[int, str]) -> dict[int, str]:
        first = self._expanded_row(table.rows[0])
        names = [self._normalize_text(text) for text, _ in first]
        if any(name.startswith("РЗ-") for name in names):
            return {index: name for index, name in enumerate(names) if name.startswith("РЗ-")}
        # The second table in the provided DOCX is a continuation without its own group header.
        schedule_cols = max(len(first) - 2, 0)
        if schedule_cols == 5:
            return {2: "РЗ-251", 3: "РЗ-252", 4: "РЗ-252", 5: "РЗ-252", 6: "РЗ-253"}
        return previous

    def _groups_for_span(self, row: list[tuple[str, int]], col_index: int, group_map: dict[int, str]) -> list[str]:
        _, cell_id = row[col_index]
        indices = [index for index, (_, other_id) in enumerate(row) if other_id == cell_id]
        groups = [group_map[index] for index in indices if index in group_map]
        return sorted(set(groups))

    def _parse_cell(self, text: str, lesson_number: int) -> dict[str, object]:
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        first = lines[0] if lines else ""
        week_start = week_end = None
        week_parity = "any"
        lesson_type: str | None = None
        subject_lines = lines[:]
        match = WEEK_PREFIX_RE.match(first)
        if match:
            week_start = int(match.group("start"))
            week_end = int(match.group("end"))
            rest = match.group("rest").strip()
            tokens = rest.split()
            if tokens and tokens[0].startswith("н/пар"):
                week_parity = "odd"
                tokens = tokens[1:]
            elif tokens and tokens[0].startswith("пар"):
                week_parity = "even"
                tokens = tokens[1:]
            if tokens and tokens[0].lower() in TYPE_WORDS:
                lesson_type = tokens[0].lower()
                tokens = tokens[1:]
            subject_lines = ([" ".join(tokens)] if tokens else []) + lines[1:]
        meeting_url = self._extract_url(text)
        cleaned_lines = [URL_RE.sub("", line).strip() for line in subject_lines]
        cleaned_lines = [line for line in cleaned_lines if line]
        teacher = None
        if cleaned_lines and self._looks_like_teacher(cleaned_lines[-1]):
            teacher = cleaned_lines.pop()
        if not lesson_type and cleaned_lines:
            first_words = cleaned_lines[0].split()
            if first_words and first_words[0].lower() in TYPE_WORDS:
                lesson_type = first_words[0].lower()
                cleaned_lines[0] = " ".join(first_words[1:]).strip()
        subject = "\n".join(line for line in cleaned_lines if line).strip()
        starts_at, ends_at = LESSON_TIMES.get(lesson_number, (time(0), time(0)))
        return {
            "starts_at": starts_at,
            "ends_at": ends_at,
            "subject": subject or text[:120],
            "teacher": teacher,
            "meeting_url": meeting_url,
            "description": text,
            "lesson_type": lesson_type,
            "week_start": week_start,
            "week_end": week_end,
            "week_parity": week_parity,
        }

    def _validate(self, result: ParseResult) -> ParseResult:
        conflicts: dict[tuple[str, int, int, str], list[str]] = defaultdict(list)
        for lesson in result.lessons:
            location = f"table {lesson.source_table}, row {lesson.source_row}, col {lesson.source_col}"
            if not lesson.subject:
                result.errors.append(f"{location}: missing subject.")
            if lesson.starts_at >= lesson.ends_at:
                result.errors.append(f"{location}: end time must be after start time.")
            if lesson.meeting_url and urlparse(lesson.meeting_url).scheme not in {"http", "https"}:
                result.errors.append(f"{location}: malformed meeting URL.")
            if lesson.week_start and lesson.week_end and lesson.week_start > lesson.week_end:
                result.errors.append(f"{location}: week range start is after end.")
            conflicts[(lesson.group_name, lesson.day_of_week, lesson.lesson_number, lesson.week_parity)].append(lesson.subject)
        for key, subjects in conflicts.items():
            if len(set(subjects)) > 1:
                result.warnings.append(f"Potential conflict for group/day/lesson/parity {key}: {sorted(set(subjects))}.")
        return result

    def _extract_url(self, text: str) -> str | None:
        match = URL_RE.search(text)
        return match.group(0).rstrip(".,)") if match else None

    def _looks_like_teacher(self, line: str) -> bool:
        lowered = line.lower()
        return any(prefix in lowered for prefix in ("доц.", "проф.", "св.", "ст."))

    def _normalize_text(self, text: str) -> str:
        return "\n".join(part.strip() for part in text.replace("\xa0", " ").splitlines()).strip()

    def _compact(self, text: str) -> str:
        return self._normalize_text(text).replace("\n", "").replace(" ", "").upper()

    def _is_group_header(self, row: list[tuple[str, int]]) -> bool:
        return any(self._normalize_text(text).startswith("РЗ-") for text, _ in row)

    def _is_elective(self, text: str) -> bool:
        return self._compact(text) == "ВИБІРКОВІ"
